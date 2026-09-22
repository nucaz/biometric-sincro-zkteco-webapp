"""
HR-facing web dashboard + admin screens over the central attendance database.

Run with:
    python -m central.webapp
Then open http://localhost:5000

First run needs an admin account:
    python -m central.create_admin <username> <password>
"""
import csv
import io
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FuturesTimeoutError
from datetime import datetime, timedelta

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from flask import (
    Flask, render_template, request, redirect, url_for, session as flask_session,
    Response, flash
)
from markupsafe import Markup
from zk import const as zk_const
from sqlalchemy import func

from attendance_system import ZKTecoAttendance, check_network_reachable
from central.db import SessionLocal, init_db
from central.models import (
    Device, AttendanceLog, EmployeeMap, Employee, User, TrustedDevice, AuditLog, RoleModule,
    Schedule, EmployeeSchedule, AttendanceException,
)
from central import device_lock, attendance_rules
from central.icons import ICONS
from central.auth import (
    login_required, admin_required, super_admin_required, verify_password, hash_password,
    ROLE_LABELS, module_required, module_enabled, MODULES, CONFIGURABLE_ROLES, DEFAULT_MODULE_ACCESS,
)
from central.sync import sync_all, sync_device, backup_device
from central.backup import export_data, import_data
from central.twofa import (
    generate_secret, qr_data_uri, verify_code, new_trusted_device_token,
    hash_token, trusted_device_expiry, TRUSTED_DEVICE_COOKIE, TRUSTED_DEVICE_DAYS,
)
from central.audit import log_event, resolve_hostname, client_ip, client_user_agent

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-only-change-me")

# Idempotent (create_all only adds missing tables) -- called at import time,
# not just under __main__, so it also runs under a WSGI server like
# gunicorn/waitress that imports `central.webapp:app` directly.
init_db()

PAGE_SIZE = 200


@app.before_request
def guard():
    if request.endpoint in ("login", "verify_2fa", "static"):
        return
    if "user_id" not in flask_session:
        return redirect(url_for("login", next=request.path))


@app.after_request
def no_cache(response):
    # This is an internal HR/admin tool with no static-asset pipeline --
    # every page is server-rendered per request, so there's no upside to
    # letting a browser (or a caching proxy) reuse an old copy, and real
    # downsides: it's how a stale copy of devices.html's sync script could
    # end up running even after the server was updated with a fix.
    response.headers["Cache-Control"] = "no-store"
    return response


def _icon(name, css_class="nav-icon"):
    """Renders one entry of central/icons.py as an inline <svg> -- kept as
    a template global (not a filter) so a page can call icon('clock')
    directly. Markup() marks it safe to avoid Jinja auto-escaping the SVG
    tags into visible text."""
    return Markup(
        f'<svg class="{css_class}" viewBox="0 0 24 24" fill="none" stroke="currentColor" '
        f'stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">'
        f'{ICONS.get(name, "")}</svg>'
    )


@app.context_processor
def inject_role():
    role = flask_session.get("role")
    return {
        "current_role": role,
        "current_role_label": ROLE_LABELS.get(role, role),
        "current_username": flask_session.get("username"),
        "can_manage": role in ("admin", "super_admin"),
        "is_super_admin": role == "super_admin",
        "module_enabled": lambda m: module_enabled(role, m),
        "icon": _icon,
        "active_endpoint": request.endpoint,
    }


# -------------------------------------------------------------- audit ----

def _audit(db, action, target=None, detail=None, username=None):
    """Append one audit_log row using the client info cached on this
    session at login time (so a reverse-DNS lookup only happens once per
    session, not on every click) -- falls back to a fresh lookup for
    events that happen outside a session, like a failed login."""
    ip = flask_session.get("client_ip")
    if ip is None:
        ip = client_ip()
        hostname = resolve_hostname(ip)
        user_agent = client_user_agent()
    else:
        hostname = flask_session.get("client_hostname")
        user_agent = flask_session.get("client_user_agent")
    log_event(
        db, action,
        username=username or flask_session.get("username"),
        target=target, detail=detail,
        ip=ip, hostname=hostname, user_agent=user_agent,
    )
    db.commit()


def _audit_context():
    """Snapshot of who's making this request, as plain values -- pass this
    into a background thread instead of calling _audit() there directly,
    since flask.session/request only exist on the request thread that
    started them, not on a worker thread spawned from it."""
    return {
        "username": flask_session.get("username"),
        "ip": flask_session.get("client_ip") or client_ip(),
        "hostname": flask_session.get("client_hostname"),
        "user_agent": flask_session.get("client_user_agent") or client_user_agent(),
    }


# ------------------------------------------------------ sync hard timeout -
# No live progress view anymore -- it went through two designs (Server-Sent
# Events, then polling a background job) and both ended up adding more
# failure modes than they were worth for what's fundamentally a "click a
# button, wait, see the result" action. Back to a plain blocking request:
# the button disables itself while the form submits, the page redirects
# back with a flash message once it's done. What's still very much needed:
# sync_device()'s own retry loop is supposed to bound how long a device can
# take (3 attempts x 20s timeout + delays, under 2 minutes) -- but a flaky
# device (seen for real against one K20) can make a pyzk socket call hang
# PAST its own configured timeout instead of raising, which no amount of
# retry logic inside sync_device() can protect against. run_with_hard_timeout
# is the actual ceiling: past HARD_SYNC_TIMEOUT_SECONDS, give up and report
# an error instead of leaving the request (and the waitress/gunicorn worker
# thread serving it) hanging forever. Python can't forcibly kill the
# underlying thread if it's truly stuck in a blocking call, so a hung one
# leaks in the background (it dies with the process) -- max_workers is
# generous specifically so a few of those can't starve out real work.
_sync_executor = ThreadPoolExecutor(max_workers=8, thread_name_prefix="sync-job")
HARD_SYNC_TIMEOUT_SECONDS = 150


def run_with_hard_timeout(fn, *args, **kwargs):
    """Call fn(*args, **kwargs) with a firm ceiling. Raises TimeoutError
    (built-in, not concurrent.futures') with a clear Spanish message if it
    doesn't return in time -- callers just need one except TimeoutError."""
    future = _sync_executor.submit(fn, *args, **kwargs)
    try:
        return future.result(timeout=HARD_SYNC_TIMEOUT_SECONDS)
    except FuturesTimeoutError:
        raise TimeoutError(
            f"Sin respuesta del dispositivo tras {HARD_SYNC_TIMEOUT_SECONDS} segundos -- "
            "probablemente un problema de red con ese equipo, no de la aplicación."
        )


def run_device_op(device_pk, device_name, fn, *args, **kwargs):
    """Like run_with_hard_timeout, but for an action that opens a
    connection to one specific physical device (Probar conexión,
    Reiniciar, Apagar, Vaciar, Usuarios). Refuses to start a second one
    while a previous attempt on the SAME device is still in flight --
    including one still leaking in the background past its own hard
    timeout, see _sync_executor above -- since a ZK terminal only accepts
    one live session at a time. Uses the same DB-backed lock as
    central.sync (see central/device_lock.py), so this also won't collide
    with a "Sincronizar"/"Descargar respaldo" mid-flight for this device --
    including one running as a separate scheduled `python -m central.sync`
    process, not just another request in this same web app."""
    device_lock.acquire(device_pk, device_name, owner="webapp")

    def guarded():
        try:
            return fn(*args, **kwargs)
        finally:
            device_lock.release(device_pk)

    future = _sync_executor.submit(guarded)
    try:
        return future.result(timeout=HARD_SYNC_TIMEOUT_SECONDS)
    except FuturesTimeoutError:
        raise TimeoutError(
            f"Sin respuesta del dispositivo tras {HARD_SYNC_TIMEOUT_SECONDS} segundos -- "
            "probablemente un problema de red con ese equipo, no de la aplicación."
        )


# ---------------------------------------------------------------- auth ----

def _finish_login(user, remember_next_url=None):
    flask_session["user_id"] = user.id
    flask_session["role"] = user.role
    flask_session["username"] = user.username
    flask_session.pop("pending_2fa_user_id", None)
    flask_session.pop("pending_2fa_next", None)
    ip = client_ip()
    flask_session["client_ip"] = ip
    # Best-effort PC name via reverse DNS -- see central/audit.py for why a
    # browser can't just report its Windows computer name directly.
    flask_session["client_hostname"] = resolve_hostname(ip)
    flask_session["client_user_agent"] = client_user_agent()


def _device_is_trusted(db, user):
    token = request.cookies.get(TRUSTED_DEVICE_COOKIE)
    if not token:
        return False
    th = hash_token(token)
    return db.query(TrustedDevice).filter(
        TrustedDevice.user_id == user.id,
        TrustedDevice.token_hash == th,
        TrustedDevice.expires_at > datetime.utcnow(),
    ).first() is not None


MAX_LOGIN_ATTEMPTS = 3


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        db = SessionLocal()
        try:
            attempted_username = request.form.get("username", "").strip()
            user = db.query(User).filter(User.username == attempted_username).first()

            if user and user.locked:
                _audit(db, "login_blocked", username=user.username, detail="cuenta bloqueada por intentos fallidos")
                flash("Esta cuenta está bloqueada por demasiados intentos fallidos. Pide a un super administrador que la desbloquee en Cuentas.")
                return render_template("login.html")

            if user and user.active and verify_password(request.form.get("password", ""), user.password_hash):
                if user.failed_login_attempts:
                    user.failed_login_attempts = 0
                    db.commit()
                next_url = request.args.get("next") or url_for("dashboard")
                if user.totp_enabled and not _device_is_trusted(db, user):
                    flask_session["pending_2fa_user_id"] = user.id
                    flask_session["pending_2fa_next"] = next_url
                    return redirect(url_for("verify_2fa"))
                _finish_login(user)
                _audit(db, "login", detail="sin 2FA (no activado o equipo recordado)")
                return redirect(next_url)

            if user and user.active:
                # Only a real, active account can be locked -- an unknown
                # username or an already-disabled account doesn't count
                # attempts, there's nothing there to protect by locking it.
                user.failed_login_attempts = (user.failed_login_attempts or 0) + 1
                if user.failed_login_attempts >= MAX_LOGIN_ATTEMPTS:
                    user.locked = True
                    db.commit()
                    _audit(db, "account_locked", username=user.username, detail=f"{MAX_LOGIN_ATTEMPTS} intentos fallidos seguidos")
                    flash(f"Cuenta bloqueada tras {MAX_LOGIN_ATTEMPTS} intentos fallidos. Pide a un super administrador que la desbloquee en Cuentas.")
                else:
                    remaining = MAX_LOGIN_ATTEMPTS - user.failed_login_attempts
                    db.commit()
                    _audit(db, "login_failed", username=user.username, detail=f"contraseña incorrecta ({remaining} intento(s) restante(s))")
                    flash(f"Usuario o contraseña incorrectos. Te queda(n) {remaining} intento(s) antes de que se bloquee la cuenta.")
            else:
                _audit(db, "login_failed", username=attempted_username or None,
                       detail="usuario inactivo" if (user and not user.active) else "usuario no encontrado")
                flash("Usuario o contraseña incorrectos.")
        finally:
            db.close()
    return render_template("login.html")


@app.route("/2fa/verify", methods=["GET", "POST"])
def verify_2fa():
    pending_id = flask_session.get("pending_2fa_user_id")
    if not pending_id:
        return redirect(url_for("login"))
    db = SessionLocal()
    try:
        user = db.get(User, pending_id)
        if not user or not user.active or not user.totp_enabled:
            flask_session.pop("pending_2fa_user_id", None)
            return redirect(url_for("login"))

        if request.method == "POST":
            code = request.form.get("code", "")
            if verify_code(user.totp_secret, code):
                next_url = flask_session.get("pending_2fa_next") or url_for("dashboard")
                _finish_login(user)
                remembered = bool(request.form.get("remember_device"))
                _audit(db, "login", detail="con 2FA" + (" + equipo recordado" if remembered else ""))
                resp = redirect(next_url)
                if remembered:
                    raw_token, token_hash = new_trusted_device_token()
                    db.add(TrustedDevice(
                        user_id=user.id, token_hash=token_hash,
                        label=request.headers.get("User-Agent", "")[:200],
                        expires_at=trusted_device_expiry(),
                    ))
                    db.commit()
                    resp.set_cookie(
                        TRUSTED_DEVICE_COOKIE, raw_token,
                        max_age=TRUSTED_DEVICE_DAYS * 24 * 3600,
                        httponly=True, samesite="Lax",
                    )
                return resp
            _audit(db, "login_2fa_failed", username=user.username)
            flash("Código incorrecto. Verifica la hora de tu teléfono y vuelve a intentar.")
        return render_template("verify_2fa.html", username=user.username)
    finally:
        db.close()


@app.route("/logout")
def logout():
    if "user_id" in flask_session:
        db = SessionLocal()
        try:
            _audit(db, "logout")
        finally:
            db.close()
    flask_session.clear()
    return redirect(url_for("login"))


# ----------------------------------------------------------- dashboard ----

# Some ZKTeco terminals auto-fill this exact placeholder ("NN-<id>") as a
# user's name when they're enrolled by fingerprint/card without anyone
# typing a name on the device keypad -- it comes from the terminal itself
# (see AttendanceLog.user_name's docstring in models.py), this app never
# generates it. Showing "NN-77672037" in a name column reads like a typo
# and makes the free-text search box less useful (searching a real
# person's name can't find them until they're linked in Empleados), so
# collapse it to the same "ID <n>" fallback already used for a genuinely
# blank name -- consistent, and immediately reads as "not linked yet"
# instead of looking like a broken name. employees_unmapped() deliberately
# keeps showing the raw "NN-xxxxx" there, since that screen's whole job is
# revealing exactly what the device reports.
_NN_PLACEHOLDER = re.compile(r"^NN-?\s*\d+$", re.IGNORECASE)


def _clean_fallback_name(user_name, device_user_id):
    if user_name and not _NN_PLACEHOLDER.match(user_name.strip()):
        return user_name
    return f"ID {device_user_id}"


def _known_sedes(db):
    """Distinct sede values already in use across devices -- offered as
    autocomplete suggestions for Employee.sede (a free-text field, not a
    foreign key: an employee can exist before their sede has a device at
    all, so this is a convenience list, never a constraint)."""
    rows = db.query(Device.sede).filter(Device.sede.isnot(None)).distinct().all()
    return sorted({r[0] for r in rows if r[0]})


def _bulk_display_names(db, rows):
    """Batch version of _display_name for a whole result set -- one query
    for every (device_pk, device_user_id) pair involved instead of one
    EmployeeMap query per raw punch. With weeks of history (many punches
    per employee per day) the old per-row lookup was the actual slow part
    of loading the dashboard/exports, not the database engine itself."""
    from sqlalchemy.orm import joinedload
    device_pks = {r.device_pk for r in rows}
    if not device_pks:
        return {}
    maps = (
        db.query(EmployeeMap)
        .options(joinedload(EmployeeMap.employee))
        .filter(EmployeeMap.device_pk.in_(device_pks))
        .all()
    )
    return {(m.device_pk, m.device_user_id): m.employee.full_name for m in maps}


def _group_by_day(rows, db):
    """Collapse raw punches into one line per (device, employee, day):
    check_in = earliest punch that day, check_out = latest (when there's
    more than one). K20-class devices often don't reliably tag punches as
    in/out, so min/max-of-the-day is the robust rule rather than trusting
    each row's own status."""
    names = _bulk_display_names(db, rows)
    buckets = {}
    for r in rows:
        key = (r.device_pk, r.device_user_id, r.timestamp.date())
        bucket = buckets.setdefault(key, {
            "device_name": r.device.name,
            "sede": r.device.sede or "-",
            "employee": names.get((r.device_pk, r.device_user_id)) or _clean_fallback_name(r.user_name, r.device_user_id),
            "device_user_id": r.device_user_id,
            "date": r.timestamp.date(),
            "timestamps": [],
        })
        bucket["timestamps"].append(r.timestamp)

    grouped = []
    for bucket in buckets.values():
        ts = sorted(bucket.pop("timestamps"))
        bucket["check_in"] = ts[0]
        bucket["check_out"] = ts[-1] if len(ts) > 1 else None
        bucket["punches"] = len(ts)
        # Entrada/Salida alone hide a lunch-break out/in (or any punch in
        # between) -- keep every time that day so both the screen and the
        # exports can show what's actually behind a "revisar" count > 2,
        # instead of forcing a trip to "Marcaciones crudas" to see them.
        bucket["all_times_str"] = ", ".join(t.strftime("%H:%M:%S") for t in ts)
        grouped.append(bucket)
    grouped.sort(key=lambda g: (g["date"], g["employee"]), reverse=True)
    return grouped


def _apply_filters(query, args):
    start = args.get("start")
    end = args.get("end")
    device_id = args.get("device")
    sede = args.get("sede")
    q = (args.get("q") or "").strip()

    if start:
        query = query.filter(AttendanceLog.timestamp >= datetime.strptime(start, "%Y-%m-%d"))
    if end:
        end_dt = datetime.strptime(end, "%Y-%m-%d") + timedelta(days=1)
        query = query.filter(AttendanceLog.timestamp < end_dt)
    if device_id:
        query = query.filter(Device.device_id == device_id)
    if sede:
        query = query.filter(Device.sede == sede)
    if q:
        like = f"%{q}%"
        # Matching only the raw device fields used to miss anyone already
        # linked in Empleados whose device still reports a placeholder
        # name (e.g. "NN-77672037") -- searching their real name found
        # nothing even though the dashboard displays it correctly. Outer
        # join so a punch with no EmployeeMap row at all still matches on
        # its own raw fields.
        query = (
            query.outerjoin(
                EmployeeMap,
                (EmployeeMap.device_pk == AttendanceLog.device_pk)
                & (EmployeeMap.device_user_id == AttendanceLog.device_user_id),
            )
            .outerjoin(Employee, Employee.id == EmployeeMap.employee_id)
            .filter(
                AttendanceLog.user_name.ilike(like)
                | AttendanceLog.device_user_id.ilike(like)
                | Employee.full_name.ilike(like)
            )
        )
    return query


@app.route("/")
@module_required("dashboard")
def dashboard():
    db = SessionLocal()
    try:
        # Default to the last 7 days when no range is given, so the first
        # load of the dashboard doesn't have to group the entire history
        # (tens of thousands of rows) on every request.
        args = request.args.to_dict()
        if not args.get("start") and not args.get("end"):
            today = datetime.now().date()
            args["start"] = (today - timedelta(days=7)).strftime("%Y-%m-%d")
            args["end"] = today.strftime("%Y-%m-%d")

        base = db.query(AttendanceLog).join(Device, AttendanceLog.device_pk == Device.id)
        filtered = _apply_filters(base, args)
        view = args.get("view", "grouped")
        page = max(1, int(args.get("page", 1)))

        if view == "grouped":
            all_rows = filtered.order_by(AttendanceLog.timestamp).all()
            grouped = _group_by_day(all_rows, db)
            total = len(grouped)
            records = grouped[(page - 1) * PAGE_SIZE: page * PAGE_SIZE]
        else:
            total = filtered.count()
            rows = (filtered.order_by(AttendanceLog.timestamp.desc())
                    .offset((page - 1) * PAGE_SIZE).limit(PAGE_SIZE).all())
            names = _bulk_display_names(db, rows)
            records = [{
                "device_name": r.device.name,
                "sede": r.device.sede or "-",
                "employee": names.get((r.device_pk, r.device_user_id)) or _clean_fallback_name(r.user_name, r.device_user_id),
                "device_user_id": r.device_user_id,
                "timestamp": r.timestamp,
                "status": r.status,
            } for r in rows]

        devices = db.query(Device).order_by(Device.name).all()
        sedes = sorted({d.sede for d in devices if d.sede})

        return render_template(
            "dashboard.html",
            records=records, devices=devices, sedes=sedes, view=view,
            total=total, page=page, page_size=PAGE_SIZE, args=args,
        )
    finally:
        db.close()


def _export_table(db, args):
    """Shared by CSV/Excel/PDF export: same filters, same two views
    (grouped/raw) as the dashboard, as plain headers + row lists."""
    base = db.query(AttendanceLog).join(Device, AttendanceLog.device_pk == Device.id)
    filtered = _apply_filters(base, args).order_by(AttendanceLog.timestamp)
    rows = filtered.all()

    if args.get("view", "grouped") == "grouped":
        headers = ["Sede", "Dispositivo", "Empleado", "ID", "Fecha", "Entrada", "Salida", "Marcaciones", "Todas las marcaciones"]
        table = [[
            g["sede"], g["device_name"], g["employee"], g["device_user_id"],
            g["date"].strftime("%Y-%m-%d"), g["check_in"].strftime("%H:%M:%S"),
            g["check_out"].strftime("%H:%M:%S") if g["check_out"] else "", g["punches"],
            g["all_times_str"],
        ] for g in _group_by_day(rows, db)]
    else:
        headers = ["Sede", "Dispositivo", "Empleado", "ID", "Fecha", "Hora", "Estado"]
        names = _bulk_display_names(db, rows)
        table = [[
            r.device.sede or "", r.device.name,
            names.get((r.device_pk, r.device_user_id)) or _clean_fallback_name(r.user_name, r.device_user_id),
            r.device_user_id, r.timestamp.strftime("%Y-%m-%d"),
            r.timestamp.strftime("%H:%M:%S"), r.status,
        ] for r in rows]
    return headers, table


def _export_subtitle(args):
    parts = []
    if args.get("start") or args.get("end"):
        parts.append(f"{args.get('start', '...')} a {args.get('end', '...')}")
    if args.get("sede"):
        parts.append(f"Sede: {args['sede']}")
    if args.get("q"):
        parts.append(f"Búsqueda: \"{args['q']}\"")
    parts.append("Resumen por día" if args.get("view", "grouped") == "grouped" else "Marcaciones crudas")
    return " · ".join(parts)


@app.route("/export.csv")
@module_required("export")
def export_csv():
    db = SessionLocal()
    try:
        headers, table = _export_table(db, request.args)
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(table)
        return Response(
            buf.getvalue(), mimetype="text/csv",
            headers={"Content-Disposition": "attachment; filename=attendance_export.csv"},
        )
    finally:
        db.close()


def _template_workbook(headers, example_row=None):
    """A blank fillable sheet for a bulk-import screen: bold header row,
    one greyed-out example row (deleted by whoever fills it in) so the
    expected format is obvious without reading the instructions text,
    and columns sized to fit. Deliberately NOT run through
    _styled_workbook() -- that one's report look (title banner, frozen
    header, stripes) fits a read-only export, not a sheet meant to be
    edited and re-uploaded."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla"

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    if example_row:
        for col, value in enumerate(example_row, start=1):
            cell = ws.cell(row=2, column=col, value=value)
            cell.font = Font(italic=True, color="9CA3AF")

    for col, title in enumerate(headers, start=1):
        letter = get_column_letter(col)
        example_len = len(str(example_row[col - 1])) if example_row and col - 1 < len(example_row) else 0
        ws.column_dimensions[letter].width = min(max(len(title), example_len) + 4, 40)

    ws.freeze_panes = "A2"
    return wb


def _styled_workbook(sheet_title, report_title, subtitle, headers, table):
    """Shared look for every Excel export in this app (header banner,
    striped rows, autofilter) -- pulled out so a new report doesn't have
    to re-copy this block, it just supplies its own headers/table."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.merge_cells("A1:{}1".format(get_column_letter(len(headers))))
    ws["A1"] = report_title
    ws["A1"].font = Font(size=14, bold=True, color="1F2937")

    ws.merge_cells("A2:{}2".format(get_column_letter(len(headers))))
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="6B7280")

    ws.merge_cells("A3:{}3".format(get_column_letter(len(headers))))
    ws["A3"] = f"Generado {datetime.now().strftime('%Y-%m-%d %H:%M')} · {len(table)} registro(s) · Develop Juan Carlos Aguirre Alvarado"
    ws["A3"].font = Font(size=8, color="9CA3AF")

    header_row = 5
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    stripe_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    for r, row in enumerate(table, start=header_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if (r - header_row) % 2 == 0:
                cell.fill = stripe_fill

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        max_len = max([len(str(headers[col - 1]))] + [len(str(row[col - 1])) for row in table]) if table else len(str(headers[col - 1]))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 40)

    ws.freeze_panes = f"A{header_row + 1}"
    if table:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(table)}"
    return wb


@app.route("/export.xlsx")
@module_required("export")
def export_xlsx():
    db = SessionLocal()
    try:
        headers, table = _export_table(db, request.args)
        wb = _styled_workbook("Marcaciones", "Asistencia - Panel RRHH", _export_subtitle(request.args), headers, table)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"attendance_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return Response(
            buf.read(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


class _ReportPDF:
    """Thin wrapper so we don't repeat the header/footer/striping dance
    between the two PDF exports (dashboard + device users, if that ever
    grows one)."""

    def __init__(self, title, subtitle, orientation="P"):
        from fpdf import FPDF

        class PDF(FPDF):
            def footer(self):
                self.set_y(-15)
                self.set_font("Arial", size=8)
                self.set_text_color(150, 150, 150)
                self.cell(0, 10, f"Página {self.page_no()}/{{nb}} · Develop Juan Carlos Aguirre Alvarado - Infraestructura TI & Ciberseguridad Apps", align="C")

        self.pdf = PDF(orientation=orientation)
        self.pdf.alias_nb_pages()
        self.pdf.add_page()
        self.pdf.set_font("Arial", "B", 16)
        self.pdf.set_text_color(31, 41, 55)
        self.pdf.cell(0, 10, title, ln=True)
        self.pdf.set_font("Arial", size=10)
        self.pdf.set_text_color(107, 114, 128)
        self.pdf.cell(0, 6, subtitle, ln=True)
        self.pdf.set_text_color(0, 0, 0)
        self.pdf.ln(4)

    def table(self, headers, rows, col_widths):
        pdf = self.pdf
        pdf.set_font("Arial", "B", 9)
        pdf.set_fill_color(31, 41, 55)
        pdf.set_text_color(255, 255, 255)
        for w, h in zip(col_widths, headers):
            pdf.cell(w, 8, str(h), border=0, fill=True, align="C")
        pdf.ln()
        pdf.set_font("Arial", size=8)
        pdf.set_text_color(20, 20, 20)
        for i, row in enumerate(rows):
            pdf.set_fill_color(243, 244, 246) if i % 2 else pdf.set_fill_color(255, 255, 255)
            for w, value in zip(col_widths, row):
                text = str(value)
                if pdf.get_string_width(text) > w - 2:
                    while pdf.get_string_width(text + "...") > w - 2 and len(text) > 1:
                        text = text[:-1]
                    text += "..."
                pdf.cell(w, 6, text, border="B", fill=True)
            pdf.ln()

    def output(self):
        return bytes(self.pdf.output())


@app.route("/export.pdf")
@module_required("export")
def export_pdf():
    db = SessionLocal()
    try:
        headers, table = _export_table(db, request.args)
        is_grouped = request.args.get("view", "grouped") == "grouped"
        # Landscape for the grouped view -- "Todas las marcaciones" can list
        # half a dozen times for one day and needs the extra width portrait
        # doesn't have.
        report = _ReportPDF("Reporte de Asistencia", _export_subtitle(request.args), orientation="L" if is_grouped else "P")
        if is_grouped:
            col_widths = [24, 22, 31, 16, 20, 18, 18, 22, 106]
        else:
            col_widths = [28, 25, 55, 22, 25, 22, 25]
        report.table(headers, table, col_widths)

        filename = f"attendance_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        return Response(
            report.output(), mimetype="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


# ------------------------------------------------------------- devices ----

@app.route("/devices")
@module_required("devices")
def devices_status():
    db = SessionLocal()
    try:
        devices = db.query(Device).filter(Device.deleted.is_(False)).order_by(Device.name).all()
        deleted_devices = db.query(Device).filter(Device.deleted.is_(True)).order_by(Device.name).all()
        deleted_log_counts = {
            device_pk: count for device_pk, count in
            db.query(AttendanceLog.device_pk, func.count(AttendanceLog.id))
            .filter(AttendanceLog.device_pk.in_([d.id for d in deleted_devices]))
            .group_by(AttendanceLog.device_pk).all()
        } if deleted_devices else {}
        return render_template(
            "devices.html", devices=devices, deleted_devices=deleted_devices,
            deleted_log_counts=deleted_log_counts,
        )
    finally:
        db.close()


@app.route("/sync", methods=["POST"])
@admin_required
def sync_now():
    """Manual trigger for the same pull that central.sync does on a
    schedule -- blocks until every active device has been polled or the
    hard timeout gives up on one (can take a while with a lot of devices
    or history; the button disables itself client-side while it waits)."""
    ctx = _audit_context()

    def work():
        results = sync_all()
        summary = "; ".join(f"{name}: {result}" for name, result in results.items())
        db = SessionLocal()
        try:
            log_event(db, "sync_all", detail=summary, **ctx)
            db.commit()
        finally:
            db.close()
        return summary

    try:
        flash(f"Sincronización completa. {run_with_hard_timeout(work)}")
    except TimeoutError as e:
        db = SessionLocal()
        try:
            log_event(db, "sync_all", detail=f"ERROR: {e}", **ctx)
            db.commit()
        finally:
            db.close()
        flash(f"Sincronización interrumpida: {e}")
    return redirect(url_for("devices_status"))


@app.route("/devices/<int:device_pk>/sync", methods=["POST"])
@admin_required
def device_sync_one(device_pk):
    """Same as 'Sincronizar todos' but for just this one device -- handy
    right after fixing a device's IP/comm key without waiting for the
    others to be polled too."""
    ctx = _audit_context()
    db = SessionLocal()
    try:
        device_name = (db.get(Device, device_pk) or abort_404()).name
    finally:
        db.close()

    def work():
        db2 = SessionLocal()
        try:
            device = db2.get(Device, device_pk)
            try:
                count = sync_device(db2, device)
                log_event(db2, "device_sync", target=device.name, detail=f"{count} registro(s) nuevo(s)", **ctx)
                db2.commit()
                return count
            except Exception as e:
                db2.rollback()
                log_event(db2, "device_sync", target=device.name, detail=f"ERROR: {e}", **ctx)
                db2.commit()
                raise
        finally:
            db2.close()

    try:
        count = run_with_hard_timeout(work)
        flash(f"'{device_name}' sincronizado: {count} registro(s) nuevo(s).")
    except TimeoutError as e:
        flash(f"'{device_name}': {e}")
    except Exception as e:
        flash(f"Error sincronizando '{device_name}': {e}")
    return redirect(url_for("devices_status"))


@app.route("/devices/<int:device_pk>/backup", methods=["POST"])
@admin_required
def device_backup(device_pk):
    """Descarga un CSV con TODAS las marcaciones que hay ahora mismo EN
    EL EQUIPO (no lo que ya está en la base central), y de paso asegura
    que la base central también las tenga (mismo dedup que 'Sincronizar').
    Pensado para usarse justo antes de 'Vaciar dispositivo': el archivo
    queda como respaldo independiente, sin depender de confiar en la base
    central."""
    ctx = _audit_context()
    db = SessionLocal()
    try:
        device_name = (db.get(Device, device_pk) or abort_404()).name
    finally:
        db.close()

    def work():
        db2 = SessionLocal()
        try:
            device = db2.get(Device, device_pk)
            try:
                raw, new_count = backup_device(db2, device)
                total = 0 if raw is None else len(raw)
                log_event(db2, "device_backup", target=device.name,
                          detail=f"{total} registro(s) en el equipo, {new_count} nuevo(s) guardado(s) en la base central", **ctx)
                db2.commit()
                return raw
            except Exception as e:
                db2.rollback()
                log_event(db2, "device_backup", target=device.name, detail=f"ERROR: {e}", **ctx)
                db2.commit()
                raise
        finally:
            db2.close()

    try:
        raw = run_with_hard_timeout(work)
    except TimeoutError as e:
        flash(f"'{device_name}': {e}")
        return redirect(url_for("devices_status"))
    except Exception as e:
        flash(f"Error generando respaldo de '{device_name}': {e}")
        return redirect(url_for("devices_status"))

    if raw is None or raw.empty:
        flash(f"'{device_name}' no tiene marcaciones guardadas en el equipo ahora mismo -- no hay nada que respaldar.")
        return redirect(url_for("devices_status"))

    buf = io.StringIO()
    raw.to_csv(buf, index=False)
    filename = f"respaldo_{device_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(
        buf.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.route("/devices/new", methods=["GET", "POST"])
@admin_required
def device_new():
    if request.method == "POST":
        db = SessionLocal()
        try:
            ip = request.form["ip"].strip()
            port = int(request.form.get("port") or 4370)
            comm_password = request.form.get("comm_password", "").strip()
            terminal_no = request.form.get("terminal_no", "").strip()
            device = Device(
                device_id=f"{ip}:{port}",  # upgraded to the real serial once "Probar conexión" runs
                name=request.form["name"].strip(),
                sede=request.form.get("sede", "").strip() or None,
                ip=ip, port=port, active=1,
                comm_password=int(comm_password) if comm_password.isdigit() else 0,
                terminal_no=int(terminal_no) if terminal_no.isdigit() else None,
            )
            db.add(device)
            db.commit()
            _audit(db, "device_add", target=device.name, detail=f"{ip}:{port}")
            flash(f"Dispositivo '{device.name}' agregado. Usa 'Probar conexión' para confirmarlo.")
            return redirect(url_for("devices_status"))
        finally:
            db.close()
    return render_template("device_form.html", device=None)


@app.route("/devices/<int:device_pk>/edit", methods=["GET", "POST"])
@admin_required
def device_edit(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        if request.method == "POST":
            comm_password = request.form.get("comm_password", "").strip()
            terminal_no = request.form.get("terminal_no", "").strip()
            device.name = request.form["name"].strip()
            device.sede = request.form.get("sede", "").strip() or None
            device.ip = request.form["ip"].strip()
            device.port = int(request.form.get("port") or 4370)
            device.comm_password = int(comm_password) if comm_password.isdigit() else 0
            device.terminal_no = int(terminal_no) if terminal_no.isdigit() else None
            device.active = 1 if request.form.get("active") else 0
            db.commit()
            _audit(db, "device_edit", target=device.name)
            flash(f"Dispositivo '{device.name}' actualizado.")
            return redirect(url_for("devices_status"))
        history_count = db.query(AttendanceLog).filter(AttendanceLog.device_pk == device.id).count()
        return render_template("device_form.html", device=device, history_count=history_count)
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/delete", methods=["POST"])
@admin_required
def device_delete(device_pk):
    """"Eliminar" never drops attendance_log/employee_map history and never
    drops the devices row itself -- it hides the device from this list and
    from sync (both manual and scheduled), same as unchecking 'Activo',
    plus taking it out of view entirely. See /devices/<id>/restore to
    bring it back."""
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        history_count = db.query(AttendanceLog).filter(AttendanceLog.device_pk == device.id).count()
        device.deleted = True
        device.active = 0
        db.commit()
        _audit(db, "device_delete", target=device.name, detail=f"oculto, {history_count} marcación(es) conservada(s) en el historial")
        flash(f"'{device.name}' eliminado de la lista. Su historial ({history_count} marcación(es)) se conserva intacto en Marcaciones -- usa 'Ver eliminados' más abajo para restaurarlo si hace falta.")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/restore", methods=["POST"])
@admin_required
def device_restore(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device.deleted = False
        db.commit()
        _audit(db, "device_restore", target=device.name)
        flash(f"'{device.name}' restaurado -- sigue inactivo (no se sincroniza) hasta que lo actives desde Editar.")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/test", methods=["POST"])
@admin_required
def device_test(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password
    finally:
        db.close()

    reachable, detail = check_network_reachable(device_ip, device_port)
    if not reachable:
        db = SessionLocal()
        try:
            _audit(db, "device_test", target=device_name, detail=f"ERROR (red): {detail}")
        finally:
            db.close()
        flash(f"'{device_name}' ({device_ip}:{device_port}) no responde en la red: {detail}")
        return redirect(url_for("devices_status"))

    def work():
        system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
        system.connect()
        if not system.conn:
            return None
        try:
            return {
                "model": system.device_info.get('model', '').strip(),
                "serial": system.device_info.get('serial_number', '').strip(),
            }
        finally:
            system.disconnect()

    db = SessionLocal()
    try:
        try:
            result = run_device_op(device_pk, device_name, work)
        except Exception as e:
            _audit(db, "device_test", target=device_name, detail=f"ERROR: {e}")
            flash(f"'{device_name}': {e}")
            return redirect(url_for("devices_status"))

        if result is None:
            hint = " Revisa la llave de comunicación (comm key) configurada en el equipo." if comm_password == 0 else " Revisa que la llave de comunicación configurada aquí coincida con la del equipo."
            _audit(db, "device_test", target=device_name, detail="ERROR (protocolo ZK rechazado)")
            flash(f"'{device_name}' responde en la red pero rechazó la conexión del protocolo ZK.{hint}")
            return redirect(url_for("devices_status"))

        device = db.get(Device, device_pk) or abort_404()
        model, serial = result["model"], result["serial"]
        device.model = model or device.model
        if serial and device.device_id != serial:
            clash = db.query(Device).filter(Device.device_id == serial, Device.id != device.id).first()
            if not clash:
                device.device_id = serial
        db.commit()
        _audit(db, "device_test", target=device.name, detail=f"OK - modelo {model or '?'}, serial {serial or '?'}")
        flash(f"Conexión OK con '{device.name}': modelo {model or '?'}, serial {serial or '?'}.")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/ping", methods=["POST"])
@admin_required
def device_ping(device_pk):
    """Pure network-reachability test (TCP connect) -- no ZK handshake, no
    comm key involved. Faster and simpler than 'Probar conexión' when all
    you want to know is "is it even on the network right now"."""
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        reachable, detail = check_network_reachable(device.ip, device.port)
        _audit(db, "device_ping", target=device.name, detail="OK" if reachable else f"ERROR: {detail}")
        if reachable:
            flash(f"Ping OK: '{device.name}' ({device.ip}:{device.port}) responde en la red.")
        else:
            flash(f"Ping falló: '{device.name}' ({device.ip}:{device.port}) — {detail}")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/restart", methods=["POST"])
@admin_required
def device_restart(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password
    finally:
        db.close()

    def work():
        system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
        system.connect()
        if not system.conn:
            raise RuntimeError(system.last_connect_error or "no se pudo conectar")
        try:
            system.restart_device()
        finally:
            system.disconnect()

    db = SessionLocal()
    try:
        try:
            run_device_op(device_pk, device_name, work)
            _audit(db, "device_restart", target=device_name)
            flash(f"Se envió el reinicio a '{device_name}'. Tardará unos segundos en volver a estar disponible.")
        except Exception as e:
            _audit(db, "device_restart", target=device_name, detail=f"ERROR: {e}")
            flash(f"Error al reiniciar '{device_name}': {e}")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/poweroff", methods=["POST"])
@admin_required
def device_poweroff(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password
    finally:
        db.close()

    def work():
        system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
        system.connect()
        if not system.conn:
            raise RuntimeError(system.last_connect_error or "no se pudo conectar")
        try:
            system.poweroff_device()
        finally:
            system.disconnect()

    db = SessionLocal()
    try:
        try:
            run_device_op(device_pk, device_name, work)
            _audit(db, "device_poweroff", target=device_name)
            flash(f"Se envió el apagado a '{device_name}'. A diferencia de reiniciar, no volverá a encender solo.")
        except Exception as e:
            _audit(db, "device_poweroff", target=device_name, detail=f"ERROR: {e}")
            flash(f"Error al apagar '{device_name}': {e}")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/clear_attendance", methods=["POST"])
@admin_required
def device_clear_attendance(device_pk):
    """Wipes the punches stored ON THE PHYSICAL TERMINAL so it starts
    accumulating from zero -- does NOT touch attendance_log in the central
    database (that history stays exactly as it is). Requires typing the
    device's name to confirm (see devices.html) since it's irreversible on
    the device side."""
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password
        if request.form.get("confirm_name", "").strip() != device_name:
            flash(f"El nombre escrito no coincide con '{device_name}' -- no se vació nada.")
            return redirect(url_for("devices_status"))
    finally:
        db.close()

    def work():
        system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
        system.connect()
        if not system.conn:
            raise RuntimeError(system.last_connect_error or "no se pudo conectar")
        try:
            system.clear_attendance_log()
        finally:
            system.disconnect()

    db = SessionLocal()
    try:
        try:
            run_device_op(device_pk, device_name, work)
            _audit(db, "device_clear_attendance", target=device_name,
                   detail="marcaciones borradas del equipo (el historial central se conserva)")
            flash(f"'{device_name}' vaciado: sus marcaciones se borraron del equipo. El historial ya sincronizado en el panel se conserva intacto.")
        except Exception as e:
            _audit(db, "device_clear_attendance", target=device_name, detail=f"ERROR: {e}")
            flash(f"Error al vaciar '{device_name}': {e}")
        return redirect(url_for("devices_status"))
    finally:
        db.close()


def abort_404():
    from flask import abort
    abort(404)


# --------------------------------------------------------- device users ----
# Employees enrolled directly on a physical terminal (fingerprint/card),
# as opposed to the Employee master record kept centrally.

PRIVILEGE_LABELS = {
    zk_const.USER_DEFAULT: "Usuario",
    zk_const.USER_ADMIN: "Administrador",
    zk_const.USER_ENROLLER: "Enrolador",
    zk_const.USER_MANAGER: "Supervisor",
}


@app.route("/devices/<int:device_pk>/users")
@module_required("device_users")
def device_users(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device_name = device.name
        device_ip, device_port, comm_password = device.ip, device.port, device.comm_password
    finally:
        db.close()

    def work():
        system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
        system.connect()
        if not system.conn:
            raise RuntimeError("no respondió")
        try:
            return system.list_device_users()
        finally:
            system.disconnect()

    try:
        users = run_device_op(device_pk, device_name, work)
    except TimeoutError as e:
        flash(f"'{device_name}': {e}")
        return redirect(url_for("devices_status"))
    except Exception as e:
        flash(f"No se pudo conectar a '{device_name}' para listar usuarios: {e}")
        return redirect(url_for("devices_status"))

    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        for u in users:
            u['privilege_label'] = PRIVILEGE_LABELS.get(u['privilege'], str(u['privilege']))
            mapped = db.query(EmployeeMap).filter(
                EmployeeMap.device_pk == device.id, EmployeeMap.device_user_id == str(u['user_id'])
            ).first()
            u['employee'] = mapped.employee if mapped else None
        return render_template("device_users.html", device=device, users=users)
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/users/new", methods=["GET", "POST"])
@admin_required
def device_user_new(device_pk):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        employees = db.query(Employee).filter(Employee.active.is_(True)).order_by(Employee.full_name).all()
        if request.method == "POST":
            user_id = request.form["user_id"].strip()
            name = request.form["name"].strip()
            privilege = zk_const.USER_ADMIN if request.form.get("privilege") == "admin" else zk_const.USER_DEFAULT
            card = request.form.get("card", "").strip()
            device_password = request.form.get("device_password", "").strip()
            employee_id = request.form.get("employee_id")
            device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password

            def work():
                system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
                system.connect()
                if not system.conn:
                    raise RuntimeError(system.last_connect_error or "no se pudo conectar")
                try:
                    return system.create_or_update_user(
                        user_id=user_id, name=name, privilege=privilege,
                        card=int(card) if card.isdigit() else 0, password=device_password,
                    )
                finally:
                    system.disconnect()

            try:
                created = run_device_op(device_pk, device_name, work)
            except Exception as e:
                flash(f"Error escribiendo en el dispositivo: {e}")
                return redirect(url_for("device_users", device_pk=device.id))

            if employee_id:
                existing_map = db.query(EmployeeMap).filter(
                    EmployeeMap.device_pk == device.id, EmployeeMap.device_user_id == user_id
                ).first()
                if existing_map:
                    existing_map.employee_id = int(employee_id)
                else:
                    db.add(EmployeeMap(device_pk=device.id, device_user_id=user_id, employee_id=int(employee_id)))
                db.commit()

            _audit(db, "device_user_add" if created else "device_user_edit",
                   target=f"{device.name}: {user_id}", detail=name)
            flash(f"Usuario '{name}' {'creado' if created else 'actualizado'} en '{device.name}'.")
            return redirect(url_for("device_users", device_pk=device.id))
        return render_template("device_user_form.html", device=device, employees=employees)
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/users/<user_id>/edit", methods=["GET", "POST"])
@admin_required
def device_user_edit(device_pk, user_id):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        employees = db.query(Employee).filter(Employee.active.is_(True)).order_by(Employee.full_name).all()
        current_map = db.query(EmployeeMap).filter(
            EmployeeMap.device_pk == device.id, EmployeeMap.device_user_id == user_id
        ).first()

        if request.method == "POST":
            name = request.form["name"].strip()
            privilege = zk_const.USER_ADMIN if request.form.get("privilege") == "admin" else zk_const.USER_DEFAULT
            card = request.form.get("card", "").strip()
            device_password = request.form.get("device_password", "").strip()
            employee_id = request.form.get("employee_id")
            device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password

            def work():
                # create_or_update_user matches by user_id and reuses the
                # existing internal uid, so this writes an update, not a
                # second duplicate person.
                system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
                system.connect()
                if not system.conn:
                    raise RuntimeError(system.last_connect_error or "no se pudo conectar")
                try:
                    system.create_or_update_user(
                        user_id=user_id, name=name, privilege=privilege,
                        card=int(card) if card.isdigit() else 0, password=device_password,
                    )
                finally:
                    system.disconnect()

            try:
                run_device_op(device_pk, device_name, work)
            except Exception as e:
                flash(f"Error escribiendo en el dispositivo: {e}")
                return redirect(url_for("device_users", device_pk=device.id))

            if employee_id:
                if current_map:
                    current_map.employee_id = int(employee_id)
                else:
                    db.add(EmployeeMap(device_pk=device.id, device_user_id=user_id, employee_id=int(employee_id)))
                db.commit()
            elif current_map:
                db.delete(current_map)
                db.commit()

            _audit(db, "device_user_edit", target=f"{device.name}: {user_id}", detail=name)
            flash(f"Usuario '{name}' actualizado en '{device.name}' y en la base central.")
            return redirect(url_for("device_users", device_pk=device.id))

        # GET: pull the user's current data live from the device to prefill the form.
        device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password

        def work():
            system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
            system.connect()
            if not system.conn:
                raise RuntimeError(system.last_connect_error or "no se pudo conectar")
            try:
                return next((u for u in system.list_device_users() if str(u['user_id']) == str(user_id)), None)
            finally:
                system.disconnect()

        try:
            live_user = run_device_op(device_pk, device_name, work)
        except Exception as e:
            flash(f"No se pudo conectar a '{device_name}' para leer los datos del usuario: {e}")
            return redirect(url_for("device_users", device_pk=device.id))
        if not live_user:
            flash(f"El usuario {user_id} ya no existe en '{device.name}'.")
            return redirect(url_for("device_users", device_pk=device.id))

        return render_template(
            "device_user_form.html", device=device, employees=employees,
            editing_user=live_user, current_employee_id=current_map.employee_id if current_map else None,
        )
    finally:
        db.close()


@app.route("/devices/<int:device_pk>/users/<user_id>/delete", methods=["POST"])
@admin_required
def device_user_delete(device_pk, user_id):
    db = SessionLocal()
    try:
        device = db.get(Device, device_pk) or abort_404()
        device_name, device_ip, device_port, comm_password = device.name, device.ip, device.port, device.comm_password
    finally:
        db.close()

    def work():
        system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
        system.connect()
        if not system.conn:
            raise RuntimeError(system.last_connect_error or "no se pudo conectar")
        try:
            system.delete_device_user(user_id)
        finally:
            system.disconnect()

    db = SessionLocal()
    try:
        try:
            run_device_op(device_pk, device_name, work)
            _audit(db, "device_user_delete", target=f"{device_name}: {user_id}")
            flash(f"Usuario {user_id} eliminado de '{device_name}' (su historial de marcaciones se conserva).")
        except Exception as e:
            _audit(db, "device_user_delete", target=f"{device_name}: {user_id}", detail=f"ERROR: {e}")
            flash(f"Error eliminando en el dispositivo: {e}")
        return redirect(url_for("device_users", device_pk=device_pk))
    finally:
        db.close()


# --------------------------------------------------------- employee master --

@app.route("/employees")
@module_required("employees")
def employees_list():
    db = SessionLocal()
    try:
        employees = db.query(Employee).order_by(Employee.full_name).all()
        return render_template("employees.html", employees=employees)
    finally:
        db.close()


def _parse_date_field(value):
    value = (value or "").strip()
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def _employee_profile_fields_from_form(form):
    """Shared by employee_new/employee_edit: the extended-profile fields,
    plus full_name computed from Nombres/Apellidos when any of those are
    filled in -- see the Employee class docstring in models.py for why
    full_name isn't simply replaced by the structured fields."""
    nombre = form.get("nombre", "").strip() or None
    apellido_paterno = form.get("apellido_paterno", "").strip() or None
    apellido_materno = form.get("apellido_materno", "").strip() or None
    computed_name = " ".join(p for p in (nombre, apellido_paterno, apellido_materno) if p)
    full_name = computed_name or form.get("full_name", "").strip()
    return {
        "full_name": full_name,
        "nombre": nombre,
        "apellido_paterno": apellido_paterno,
        "apellido_materno": apellido_materno,
        "fecha_nacimiento": _parse_date_field(form.get("fecha_nacimiento")),
        "celular": form.get("celular", "").strip() or None,
        "correo_personal": form.get("correo_personal", "").strip() or None,
        "correo_corporativo": form.get("correo_corporativo", "").strip() or None,
        "direccion": form.get("direccion", "").strip() or None,
        "area": form.get("area", "").strip() or None,
        "sede": form.get("sede", "").strip() or None,
    }


@app.route("/employees/new", methods=["GET", "POST"])
@admin_required
def employee_new():
    if request.method == "POST":
        db = SessionLocal()
        try:
            fields = _employee_profile_fields_from_form(request.form)
            db.add(Employee(
                code=request.form.get("code", "").strip() or None,
                active=True,
                **fields,
            ))
            db.commit()
            _audit(db, "employee_add", target=fields["full_name"])
            flash("Empleado agregado al maestro.")
            return redirect(url_for("employees_list"))
        finally:
            db.close()
    db = SessionLocal()
    try:
        return render_template("employee_form.html", employee=None, known_sedes=_known_sedes(db))
    finally:
        db.close()


CODE_COLUMN_NAMES = {"codigo", "código", "code", "dni", "id"}
# Deliberately just "nombre" here, not "nombre completo" -- a file using
# the OLD single-column convention (one full name per "nombre" cell) and
# a file using the NEW structured convention (given name only, alongside
# separate apellido paterno/materno columns) both read a bare "nombre"
# header the same way in _employee_row_fields(): computed = nombre +
# apellido_paterno + apellido_materno joined, which is just `nombre`
# unchanged when the apellido columns aren't present. No mode flag needed.
GIVEN_NAME_COLUMN_NAMES = {"nombre"}
FULL_NAME_COLUMN_NAMES = {"nombre completo", "full_name", "name", "empleado"}
APELLIDO_PATERNO_COLUMN_NAMES = {"apellido paterno", "apellido_paterno", "paterno"}
APELLIDO_MATERNO_COLUMN_NAMES = {"apellido materno", "apellido_materno", "materno"}
FECHA_NACIMIENTO_COLUMN_NAMES = {"fecha de nacimiento", "fecha_nacimiento", "nacimiento"}
CELULAR_COLUMN_NAMES = {"celular", "telefono", "teléfono", "movil", "móvil"}
CORREO_PERSONAL_COLUMN_NAMES = {"correo personal", "correo_personal", "email personal"}
CORREO_CORPORATIVO_COLUMN_NAMES = {"correo corporativo", "correo_corporativo", "email corporativo", "correo"}
DIRECCION_COLUMN_NAMES = {"direccion", "dirección"}
AREA_COLUMN_NAMES = {"area", "área"}
SEDE_COLUMN_NAMES = {"sede"}

# The full plantilla, in the order RRHH gave it -- used both to build the
# downloadable template (employees_import_template) and, implicitly,
# documented in employees_import.html.
EMPLOYEE_TEMPLATE_HEADERS = [
    "dni", "nombre", "apellido paterno", "apellido materno",
    "fecha de nacimiento", "celular", "correo personal", "correo corporativo",
    "direccion", "area", "sede",
]
EMPLOYEE_TEMPLATE_EXAMPLE_ROW = [
    "12345678", "Juan", "Pérez", "García",
    "1990-05-20", "999888777", "juan.perez@gmail.com", "jperez@empresa.com",
    "Av. Ejemplo 123, Pueblo Libre", "Contabilidad", "Pueblo Libre",
]


def _find_col(header, names):
    return next((i for i, h in enumerate(header) if h in names), None)


def _cell(row, col):
    if col is None or col >= len(row) or row[col] is None:
        return None
    value = row[col]
    return value.strip() if isinstance(value, str) else value


def _employee_row_fields(row, cols):
    """Parses one Excel row into Employee field values using the column
    positions found in the header -- shared shape for create and update
    (the caller decides whether a None means 'leave as-is' or 'blank')."""
    nombre = _cell(row, cols["nombre"])
    apellido_paterno = _cell(row, cols["apellido_paterno"])
    apellido_materno = _cell(row, cols["apellido_materno"])
    computed_name = " ".join(str(p) for p in (nombre, apellido_paterno, apellido_materno) if p) or None
    full_name = computed_name or _cell(row, cols["full_name"])

    fecha_raw = _cell(row, cols["fecha_nacimiento"])
    if isinstance(fecha_raw, datetime):
        fecha_nacimiento = fecha_raw.date()
    elif fecha_raw:
        fecha_nacimiento = _parse_date_field(str(fecha_raw))
    else:
        fecha_nacimiento = None

    return {
        "full_name": full_name,
        "nombre": str(nombre) if nombre else None,
        "apellido_paterno": str(apellido_paterno) if apellido_paterno else None,
        "apellido_materno": str(apellido_materno) if apellido_materno else None,
        "fecha_nacimiento": fecha_nacimiento,
        "celular": str(_cell(row, cols["celular"]) or "") or None,
        "correo_personal": str(_cell(row, cols["correo_personal"]) or "") or None,
        "correo_corporativo": str(_cell(row, cols["correo_corporativo"]) or "") or None,
        "direccion": str(_cell(row, cols["direccion"]) or "") or None,
        "area": str(_cell(row, cols["area"]) or "") or None,
        "sede": str(_cell(row, cols["sede"]) or "") or None,
    }


@app.route("/employees/import", methods=["GET", "POST"])
@admin_required
def employees_import():
    if request.method == "POST":
        file = request.files.get("xlsx_file")
        if not file or not file.filename:
            flash("Selecciona un archivo .xlsx.")
            return redirect(url_for("employees_import"))
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file.stream, read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
        except Exception as e:
            flash(f"No se pudo leer el archivo: {e}")
            return redirect(url_for("employees_import"))

        cols = {
            "code": _find_col(header, CODE_COLUMN_NAMES),
            "nombre": _find_col(header, GIVEN_NAME_COLUMN_NAMES),
            "full_name": _find_col(header, FULL_NAME_COLUMN_NAMES),
            "apellido_paterno": _find_col(header, APELLIDO_PATERNO_COLUMN_NAMES),
            "apellido_materno": _find_col(header, APELLIDO_MATERNO_COLUMN_NAMES),
            "fecha_nacimiento": _find_col(header, FECHA_NACIMIENTO_COLUMN_NAMES),
            "celular": _find_col(header, CELULAR_COLUMN_NAMES),
            "correo_personal": _find_col(header, CORREO_PERSONAL_COLUMN_NAMES),
            "correo_corporativo": _find_col(header, CORREO_CORPORATIVO_COLUMN_NAMES),
            "direccion": _find_col(header, DIRECCION_COLUMN_NAMES),
            "area": _find_col(header, AREA_COLUMN_NAMES),
            "sede": _find_col(header, SEDE_COLUMN_NAMES),
        }
        if cols["nombre"] is None and cols["full_name"] is None:
            flash(f"No encontré una columna de nombre. Encabezados vistos: {', '.join(header)}. Usa una columna llamada 'nombre' o 'nombre completo'.")
            return redirect(url_for("employees_import"))

        db = SessionLocal()
        try:
            created, updated, skipped = 0, 0, 0
            for row in rows_iter:
                if row is None or all(v is None for v in row):
                    continue
                fields = _employee_row_fields(row, cols)
                code = _cell(row, cols["code"])
                code = str(code) if code else None
                if not fields["full_name"]:
                    skipped += 1
                    continue
                existing = db.query(Employee).filter(Employee.code == code).first() if code else None
                if existing:
                    # Only overwrite a field the file actually provided a
                    # value for -- an unrelated blank cell (e.g. no
                    # "celular" column at all) must never wipe out data
                    # already on file from a previous import or manual edit.
                    for key, value in fields.items():
                        if value is not None:
                            setattr(existing, key, value)
                    updated += 1
                else:
                    db.add(Employee(code=code, active=True, **fields))
                    created += 1
            db.commit()
            _audit(db, "employee_import", detail=f"{created} creado(s), {updated} actualizado(s), {skipped} omitida(s)")
            flash(f"Importación completa: {created} creado(s), {updated} actualizado(s), {skipped} fila(s) sin nombre omitida(s).")
            return redirect(url_for("employees_list"))
        finally:
            db.close()
    return render_template("employees_import.html")


@app.route("/employees/import/template.xlsx")
@admin_required
def employees_import_template():
    wb = _template_workbook(EMPLOYEE_TEMPLATE_HEADERS, EMPLOYEE_TEMPLATE_EXAMPLE_ROW)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_empleados.xlsx"},
    )


@app.route("/employees/<int:employee_id>/edit", methods=["GET", "POST"])
@admin_required
def employee_edit(employee_id):
    db = SessionLocal()
    try:
        employee = db.get(Employee, employee_id) or abort_404()
        if request.method == "POST":
            fields = _employee_profile_fields_from_form(request.form)
            new_name = fields["full_name"]
            name_changed = new_name != employee.full_name
            was_active = employee.active
            push_to_devices = bool(request.form.get("push_to_devices"))

            employee.code = request.form.get("code", "").strip() or None
            for key, value in fields.items():
                setattr(employee, key, value)
            employee.active = bool(request.form.get("active"))
            db.commit()
            became_inactive = was_active and not employee.active

            if push_to_devices and (name_changed or became_inactive):
                results = []
                targets = []
                for link in db.query(EmployeeMap).filter(EmployeeMap.employee_id == employee.id).all():
                    device = db.get(Device, link.device_pk)
                    targets.append((device.id, device.name, device.ip, device.port, device.comm_password, link.device_user_id))

                for target_device_pk, device_name, device_ip, device_port, comm_password, device_user_id in targets:
                    if became_inactive:
                        # Deactivating an employee revokes their physical
                        # access -- delete their enrollment (fingerprint/card)
                        # from the device, same as the "Eliminar" button on
                        # device_users. Their attendance history is untouched
                        # either way (never lives on the device, and
                        # delete_device_user never touches attendance_log).
                        def work():
                            system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
                            system.connect()
                            if not system.conn:
                                raise RuntimeError(system.last_connect_error or "no se pudo conectar")
                            try:
                                system.delete_device_user(device_user_id)
                            finally:
                                system.disconnect()
                        action_label = "eliminado del equipo (desactivado)"
                    else:
                        def work():
                            system = ZKTecoAttendance(device_ip, port=device_port, password=comm_password or 0)
                            system.connect()
                            if not system.conn:
                                raise RuntimeError(system.last_connect_error or "no se pudo conectar")
                            try:
                                system.create_or_update_user(user_id=device_user_id, name=new_name)
                            finally:
                                system.disconnect()
                        action_label = "actualizado"
                    try:
                        run_device_op(target_device_pk, device_name, work)
                        results.append(f"{device_name}: {action_label}")
                    except Exception as e:
                        results.append(f"{device_name}: error ({e})")
                _audit(db, "employee_edit", target=new_name, detail="propagado a relojes: " + "; ".join(results))
                flash("Empleado actualizado. Relojes: " + (", ".join(results) if results else "no tenía dispositivos vinculados") + ".")
            else:
                _audit(db, "employee_edit", target=new_name)
                flash("Empleado actualizado (solo en la base central).")
            return redirect(url_for("employees_list"))
        schedules = db.query(Schedule).filter(Schedule.active.is_(True)).order_by(Schedule.name).all()
        current_schedule = next(
            (a for a in employee.schedule_assignments if a.effective_to is None), None
        )
        device_links = [
            {"device": db.get(Device, link.device_pk), "device_user_id": link.device_user_id}
            for link in employee.device_links
        ]
        return render_template(
            "employee_form.html", employee=employee, schedules=schedules,
            current_schedule=current_schedule, schedule_history=employee.schedule_assignments,
            device_links=device_links, known_sedes=_known_sedes(db),
        )
    finally:
        db.close()


@app.route("/employees/unmapped")
@module_required("employees_unmapped")
def employees_unmapped():
    """(device, device_user_id) pairs seen in attendance_log that have no
    EmployeeMap yet -- the quickest way to clear out "NN-xxxxx" names."""
    db = SessionLocal()
    try:
        seen = db.query(
            AttendanceLog.device_pk, AttendanceLog.device_user_id, AttendanceLog.user_name
        ).distinct().all()
        mapped = {(m.device_pk, m.device_user_id) for m in db.query(EmployeeMap).all()}
        devices = {d.id: d for d in db.query(Device).all()}
        employees = db.query(Employee).filter(Employee.active.is_(True)).order_by(Employee.full_name).all()
        # The device_user_id is often literally the employee's DNI/código --
        # if a matching Employee.code already exists, suggest it instead of
        # defaulting every row to "create a new employee".
        by_code = {e.code: e for e in employees if e.code}
        unmapped = [{
            "device_pk": pk, "device_name": devices[pk].name if pk in devices else "?",
            "device_user_id": uid, "reported_name": name,
            "suggested": by_code.get(uid),
        } for pk, uid, name in seen if (pk, uid) not in mapped]
        return render_template("employees_unmapped.html", rows=unmapped, employees=employees)
    finally:
        db.close()


@app.route("/employees/map", methods=["POST"])
@admin_required
def employee_map_create():
    db = SessionLocal()
    try:
        device_pk = int(request.form["device_pk"])
        device_user_id = request.form["device_user_id"]
        employee_id = request.form.get("employee_id")
        new_name = request.form.get("new_employee_name", "").strip()

        if not employee_id and new_name:
            employee = Employee(full_name=new_name, active=True)
            db.add(employee)
            db.flush()
            employee_id = employee.id

        if not employee_id:
            flash("Selecciona un empleado existente o escribe un nombre nuevo.")
            return redirect(url_for("employees_unmapped"))

        existing = db.query(EmployeeMap).filter(
            EmployeeMap.device_pk == device_pk, EmployeeMap.device_user_id == device_user_id
        ).first()
        if existing:
            existing.employee_id = int(employee_id)
        else:
            db.add(EmployeeMap(device_pk=device_pk, device_user_id=device_user_id, employee_id=int(employee_id)))
        db.commit()
        _audit(db, "employee_link", target=f"device {device_pk}: {device_user_id}", detail=f"employee_id={employee_id}")
        flash("Vinculado correctamente.")
        return redirect(url_for("employees_unmapped"))
    finally:
        db.close()


# ----------------------------------------------------------- horarios ----

WEEKDAY_FORM_FIELDS = [
    ("work_mon", "Lunes"), ("work_tue", "Martes"), ("work_wed", "Miércoles"),
    ("work_thu", "Jueves"), ("work_fri", "Viernes"), ("work_sat", "Sábado"), ("work_sun", "Domingo"),
]


def _parse_time_field(value, field_label):
    value = (value or "").strip()
    if not value:
        raise ValueError(f"'{field_label}' es obligatorio.")
    try:
        return datetime.strptime(value, "%H:%M").time()
    except ValueError:
        raise ValueError(f"'{field_label}' debe tener formato HH:MM.")


def _schedule_fields_from_form(form):
    name = form.get("name", "").strip()
    if not name:
        raise ValueError("El nombre del horario es obligatorio.")
    entry_time = _parse_time_field(form.get("entry_time"), "Hora de ingreso")
    exit_time = _parse_time_field(form.get("exit_time"), "Hora de salida")
    if exit_time <= entry_time:
        # Same-day shift only -- a horario nocturno que cruza medianoche
        # necesitaría comparar contra el día siguiente, no soportado hoy.
        raise ValueError("La hora de salida debe ser posterior a la hora de ingreso (no se soportan horarios que cruzan la medianoche).")
    try:
        tolerance_minutes = int(form.get("tolerance_minutes") or 0)
    except ValueError:
        raise ValueError("La tolerancia debe ser un número de minutos.")
    lunch_start_raw = (form.get("lunch_start") or "").strip()
    lunch_end_raw = (form.get("lunch_end") or "").strip()
    fields = {
        "name": name,
        "entry_time": entry_time,
        "exit_time": exit_time,
        "tolerance_minutes": tolerance_minutes,
        "lunch_start": _parse_time_field(lunch_start_raw, "Inicio de refrigerio") if lunch_start_raw else None,
        "lunch_end": _parse_time_field(lunch_end_raw, "Fin de refrigerio") if lunch_end_raw else None,
    }
    for key, _ in WEEKDAY_FORM_FIELDS:
        fields[key] = bool(form.get(key))
    return fields


@app.route("/schedules")
@module_required("schedules")
def schedules_list():
    db = SessionLocal()
    try:
        schedules = db.query(Schedule).order_by(Schedule.active.desc(), Schedule.name).all()
        return render_template("schedules.html", schedules=schedules, weekday_fields=WEEKDAY_FORM_FIELDS)
    finally:
        db.close()


@app.route("/schedules/new", methods=["GET", "POST"])
@admin_required
def schedule_new():
    if request.method == "POST":
        try:
            fields = _schedule_fields_from_form(request.form)
        except ValueError as e:
            flash(str(e))
            return render_template("schedule_form.html", schedule=None, weekday_fields=WEEKDAY_FORM_FIELDS)
        db = SessionLocal()
        try:
            db.add(Schedule(active=True, **fields))
            db.commit()
            _audit(db, "schedule_add", target=fields["name"])
            flash("Horario creado.")
            return redirect(url_for("schedules_list"))
        finally:
            db.close()
    return render_template("schedule_form.html", schedule=None, weekday_fields=WEEKDAY_FORM_FIELDS)


@app.route("/schedules/<int:schedule_id>/edit", methods=["GET", "POST"])
@admin_required
def schedule_edit(schedule_id):
    db = SessionLocal()
    try:
        schedule = db.get(Schedule, schedule_id) or abort_404()
        if request.method == "POST":
            try:
                fields = _schedule_fields_from_form(request.form)
            except ValueError as e:
                flash(str(e))
                return render_template("schedule_form.html", schedule=schedule, weekday_fields=WEEKDAY_FORM_FIELDS)
            for key, value in fields.items():
                setattr(schedule, key, value)
            schedule.active = bool(request.form.get("active"))
            db.commit()
            _audit(db, "schedule_edit", target=schedule.name)
            flash("Horario actualizado.")
            return redirect(url_for("schedules_list"))
        return render_template("schedule_form.html", schedule=schedule, weekday_fields=WEEKDAY_FORM_FIELDS)
    finally:
        db.close()


def _assign_schedule(db, employee_id, schedule_id, effective_from):
    """Shared by employee_schedule_assign (one employee, from the form) and
    schedules_import (many employees, from an Excel row each) -- closes
    whatever assignment was open before `effective_from`, so assignments
    never overlap (see EmployeeSchedule in models.py). Returns (ok, msg)
    instead of raising/flashing directly since the bulk import needs to
    keep going after one row fails, not redirect."""
    current = (
        db.query(EmployeeSchedule)
        .filter(EmployeeSchedule.employee_id == employee_id, EmployeeSchedule.effective_to.is_(None))
        .first()
    )
    if current:
        if current.effective_from >= effective_from:
            return False, "la fecha de vigencia debe ser posterior al inicio del horario que ya tiene"
        current.effective_to = effective_from - timedelta(days=1)
    db.add(EmployeeSchedule(employee_id=employee_id, schedule_id=schedule_id, effective_from=effective_from))
    return True, "ok"


@app.route("/employees/<int:employee_id>/schedule", methods=["POST"])
@admin_required
def employee_schedule_assign(employee_id):
    """Assigns a schedule to an employee starting on a given date -- see
    _assign_schedule() for the overlap-closing rule."""
    db = SessionLocal()
    try:
        employee = db.get(Employee, employee_id) or abort_404()
        schedule_id = request.form.get("schedule_id")
        if not schedule_id:
            flash("Selecciona un horario.")
            return redirect(url_for("employee_edit", employee_id=employee_id))
        schedule = db.get(Schedule, int(schedule_id)) or abort_404()

        effective_from_raw = (request.form.get("effective_from") or "").strip()
        try:
            effective_from = (
                datetime.strptime(effective_from_raw, "%Y-%m-%d").date()
                if effective_from_raw else datetime.now().date()
            )
        except ValueError:
            flash("Fecha de vigencia inválida.")
            return redirect(url_for("employee_edit", employee_id=employee_id))

        ok, msg = _assign_schedule(db, employee_id, schedule.id, effective_from)
        if not ok:
            flash(msg.capitalize() + ".")
            return redirect(url_for("employee_edit", employee_id=employee_id))

        db.commit()
        _audit(db, "employee_schedule_assign", target=employee.full_name, detail=f"horario={schedule.name} desde={effective_from}")
        flash(f"Horario '{schedule.name}' asignado a {employee.full_name} desde {effective_from.strftime('%Y-%m-%d')}.")
        return redirect(url_for("employee_edit", employee_id=employee_id))
    finally:
        db.close()


@app.route("/schedules/import", methods=["GET", "POST"])
@admin_required
def schedules_import():
    """Bulk-assign schedules from an Excel file: columns 'codigo' (matches
    Employee.code) and 'horario' (matches Schedule.name exactly, case
    insensitive), plus an optional 'vigente_desde' (YYYY-MM-DD, defaults to
    today). Mirrors employees_import()'s report-what-happened-per-row
    style rather than failing the whole file on one bad row."""
    if request.method == "POST":
        file = request.files.get("xlsx_file")
        if not file or not file.filename:
            flash("Selecciona un archivo .xlsx.")
            return redirect(url_for("schedules_import"))
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file.stream, read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
        except Exception as e:
            flash(f"No se pudo leer el archivo: {e}")
            return redirect(url_for("schedules_import"))

        code_col = next((i for i, h in enumerate(header) if h in {"codigo", "código", "code", "dni", "id"}), None)
        schedule_col = next((i for i, h in enumerate(header) if h in {"horario", "schedule"}), None)
        date_col = next((i for i, h in enumerate(header) if h in {"vigente_desde", "vigente desde", "desde"}), None)
        if code_col is None or schedule_col is None:
            flash(f"Necesito columnas 'codigo' y 'horario'. Encabezados vistos: {', '.join(header)}.")
            return redirect(url_for("schedules_import"))

        db = SessionLocal()
        try:
            schedules_by_name = {s.name.strip().lower(): s for s in db.query(Schedule).all()}
            assigned, errors = 0, []
            for row_num, row in enumerate(rows_iter, start=2):
                if row is None or all(v is None for v in row):
                    continue
                code = str(row[code_col]).strip() if code_col < len(row) and row[code_col] is not None else ""
                schedule_name = str(row[schedule_col]).strip() if schedule_col < len(row) and row[schedule_col] is not None else ""
                if not code or not schedule_name:
                    errors.append(f"Fila {row_num}: falta código u horario.")
                    continue

                employee = db.query(Employee).filter(Employee.code == code).first()
                if not employee:
                    errors.append(f"Fila {row_num}: no existe ningún empleado con código '{code}'.")
                    continue
                schedule = schedules_by_name.get(schedule_name.strip().lower())
                if not schedule:
                    errors.append(f"Fila {row_num}: no existe el horario '{schedule_name}'.")
                    continue

                effective_from = datetime.now().date()
                if date_col is not None and date_col < len(row) and row[date_col] is not None:
                    raw = row[date_col]
                    try:
                        effective_from = raw.date() if hasattr(raw, "date") else datetime.strptime(str(raw).strip(), "%Y-%m-%d").date()
                    except ValueError:
                        errors.append(f"Fila {row_num}: fecha '{raw}' inválida, se usó hoy.")

                ok, msg = _assign_schedule(db, employee.id, schedule.id, effective_from)
                if not ok:
                    errors.append(f"Fila {row_num} ({employee.full_name}): {msg}.")
                    continue
                assigned += 1

            db.commit()
            _audit(db, "schedules_import", detail=f"{assigned} asignado(s), {len(errors)} error(es)")
            flash(f"Importación completa: {assigned} horario(s) asignado(s), {len(errors)} fila(s) con error.")
            if errors:
                flash("Detalle de errores: " + " | ".join(errors[:20]) + (" ..." if len(errors) > 20 else ""))
            return redirect(url_for("schedules_list"))
        finally:
            db.close()
    return render_template("schedules_import.html")


SCHEDULE_TEMPLATE_HEADERS = ["codigo", "horario", "vigente_desde"]
SCHEDULE_TEMPLATE_EXAMPLE_ROW = ["12345678", "Turno Mañana L-S", "2026-09-01"]


@app.route("/schedules/import/template.xlsx")
@admin_required
def schedules_import_template():
    wb = _template_workbook(SCHEDULE_TEMPLATE_HEADERS, SCHEDULE_TEMPLATE_EXAMPLE_ROW)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_horarios.xlsx"},
    )


# ------------------------------------------------------ vacaciones/permisos --

@app.route("/exceptions")
@module_required("attendance_exceptions")
def exceptions_list():
    db = SessionLocal()
    try:
        employee_id = request.args.get("employee_id", type=int)
        query = db.query(AttendanceException).order_by(AttendanceException.date_from.desc())
        if employee_id:
            query = query.filter(AttendanceException.employee_id == employee_id)
        rows = query.all()
        employees = db.query(Employee).filter(Employee.active.is_(True)).order_by(Employee.full_name).all()
        return render_template(
            "exceptions_list.html", rows=rows, employees=employees,
            employee_id=employee_id, kind_labels=attendance_rules.EXCEPTION_LABELS,
        )
    finally:
        db.close()


@app.route("/exceptions/new", methods=["GET", "POST"])
@admin_required
def exception_new():
    db = SessionLocal()
    try:
        if request.method == "POST":
            employee_id = request.form.get("employee_id", type=int)
            kind = request.form.get("kind")
            date_from_raw = (request.form.get("date_from") or "").strip()
            date_to_raw = (request.form.get("date_to") or "").strip()
            note = request.form.get("note", "").strip() or None

            employee = db.get(Employee, employee_id) if employee_id else None
            if not employee:
                flash("Selecciona un empleado.")
                return redirect(url_for("exception_new"))
            if kind not in AttendanceException.KINDS:
                flash("Tipo de excepción inválido.")
                return redirect(url_for("exception_new"))
            try:
                date_from = datetime.strptime(date_from_raw, "%Y-%m-%d").date()
                date_to = datetime.strptime(date_to_raw, "%Y-%m-%d").date()
            except ValueError:
                flash("Fechas inválidas.")
                return redirect(url_for("exception_new"))
            if date_to < date_from:
                flash("'Hasta' debe ser igual o posterior a 'Desde'.")
                return redirect(url_for("exception_new"))

            db.add(AttendanceException(employee_id=employee.id, date_from=date_from, date_to=date_to, kind=kind, note=note))
            db.commit()
            _audit(db, "exception_add", target=employee.full_name, detail=f"{kind} {date_from}..{date_to}")
            flash(f"Registrado: {attendance_rules.EXCEPTION_LABELS[kind]} para {employee.full_name} del {date_from} al {date_to}.")
            return redirect(url_for("exceptions_list"))

        employees = db.query(Employee).filter(Employee.active.is_(True)).order_by(Employee.full_name).all()
        preselected = request.args.get("employee_id", type=int)
        return render_template(
            "exception_form.html", employees=employees, preselected=preselected,
            kinds=AttendanceException.KINDS, kind_labels=attendance_rules.EXCEPTION_LABELS,
        )
    finally:
        db.close()


@app.route("/exceptions/<int:exception_id>/delete", methods=["POST"])
@admin_required
def exception_delete(exception_id):
    db = SessionLocal()
    try:
        exc = db.get(AttendanceException, exception_id) or abort_404()
        target = f"{exc.employee.full_name}: {exc.kind} {exc.date_from}..{exc.date_to}"
        db.delete(exc)
        db.commit()
        _audit(db, "exception_delete", target=target)
        flash("Registro eliminado.")
        return redirect(url_for("exceptions_list"))
    finally:
        db.close()


# ------------------------------------------------------------------- panel --

@app.route("/reports/overview")
@module_required("reports_overview")
def reports_overview():
    """Consolidated KPI view: today's headcount (presentes/tardanzas/
    faltas/sin horario) plus who has piled up 2+ faltas in the last 30
    days. Deliberately its own page rather than folded into the
    Marcaciones dashboard -- that one's the post-login landing page and
    gets hit far more often, and this recomputes attendance for every
    active employee over up to 30 days each time (batched queries, not
    N+1, but still real work no other page needs on every request)."""
    db = SessionLocal()
    try:
        today = datetime.now().date()
        sede = request.args.get("sede") or None
        employees = attendance_rules.employees_in_scope(db, sede=sede)

        today_result = attendance_rules.compute_attendance(db, employees, today, today)
        counts = {"a_tiempo": 0, "tardanza": 0, "falta": 0, "justificado": 0, "descanso": 0, "sin_horario": 0}
        for _employee, days in today_result.items():
            counts[days[0].status] = counts.get(days[0].status, 0) + 1

        window_start = today - timedelta(days=29)
        window_result = attendance_rules.compute_attendance(db, employees, window_start, today)
        top_ausentismo = []
        for employee, days in window_result.items():
            n = sum(1 for d in days if d.status == "falta")
            if n >= 2:
                top_ausentismo.append((employee, n))
        top_ausentismo.sort(key=lambda x: x[1], reverse=True)

        devices = db.query(Device).order_by(Device.name).all()
        sedes = sorted({d.sede for d in devices if d.sede})

        return render_template(
            "report_overview.html",
            today=today.strftime("%Y-%m-%d"), total_employees=len(employees), counts=counts,
            top_ausentismo=top_ausentismo[:15], window_start=window_start.strftime("%Y-%m-%d"),
            sedes=sedes, sede=sede,
        )
    finally:
        db.close()


# ------------------------------------------------------ reporte de faltas --

def _absences_report_params(args):
    today = datetime.now().date()
    try:
        start_date = datetime.strptime(args.get("start") or today.strftime("%Y-%m-%d"), "%Y-%m-%d").date()
        end_date = datetime.strptime(args.get("end") or today.strftime("%Y-%m-%d"), "%Y-%m-%d").date()
    except ValueError:
        start_date = end_date = today
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    # Bounds how many (employee x day) combinations one request evaluates --
    # same spirit as HARD_SYNC_TIMEOUT_SECONDS elsewhere in this file, just
    # for a slow request instead of a hung device connection.
    if (end_date - start_date).days > 186:
        end_date = start_date + timedelta(days=186)
    sede = args.get("sede") or None
    try:
        min_faltas = max(1, int(args.get("min_faltas") or 1))
    except ValueError:
        min_faltas = 1
    return start_date, end_date, sede, min_faltas


def _absences_report_data(db, args):
    start_date, end_date, sede, min_faltas = _absences_report_params(args)
    employees = attendance_rules.employees_in_scope(db, sede=sede)
    attendance = attendance_rules.compute_attendance(db, employees, start_date, end_date)

    rows = []
    no_schedule = []
    for employee, days in attendance.items():
        fault_days = [d for d in days if d.status == "falta"]
        if any(d.status == "sin_horario" for d in days):
            no_schedule.append(employee)
        if len(fault_days) >= min_faltas:
            justified_days = [d for d in days if d.status == "justificado"]
            rows.append({
                "employee": employee,
                "faltas": len(fault_days),
                "fechas": ", ".join(d.date.strftime("%Y-%m-%d") for d in fault_days),
                "justificados": len(justified_days),
            })
    rows.sort(key=lambda r: r["faltas"], reverse=True)
    no_schedule.sort(key=lambda e: e.full_name)
    return rows, no_schedule, start_date, end_date, sede, min_faltas


@app.route("/reports/absences")
@module_required("reports_absences")
def report_absences():
    db = SessionLocal()
    try:
        rows, no_schedule, start_date, end_date, sede, min_faltas = _absences_report_data(db, request.args)
        devices = db.query(Device).order_by(Device.name).all()
        sedes = sorted({d.sede for d in devices if d.sede})
        return render_template(
            "report_absences.html",
            rows=rows, no_schedule=no_schedule, sedes=sedes,
            start=start_date.strftime("%Y-%m-%d"), end=end_date.strftime("%Y-%m-%d"),
            sede=sede, min_faltas=min_faltas,
        )
    finally:
        db.close()


@app.route("/reports/absences/export.xlsx")
@module_required("reports_absences")
def report_absences_export():
    db = SessionLocal()
    try:
        rows, no_schedule, start_date, end_date, sede, min_faltas = _absences_report_data(db, request.args)
        headers = ["Código", "Empleado", "N° de faltas", "Fechas", "N° justificados (vacaciones/permiso/descanso médico)"]
        table = [[r["employee"].code or "", r["employee"].full_name, r["faltas"], r["fechas"], r["justificados"]] for r in rows]
        subtitle = f"{start_date.strftime('%Y-%m-%d')} a {end_date.strftime('%Y-%m-%d')}"
        if sede:
            subtitle += f" · Sede: {sede}"
        subtitle += f" · Mínimo {min_faltas} falta(s)"
        wb = _styled_workbook("Faltas", "Reporte de Faltas", subtitle, headers, table)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"reporte_faltas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return Response(
            buf.read(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


# ------------------------------------------------ reporte de horas/extra --

def _hours_report_params(args):
    """Same date/sede parsing as _absences_report_params -- kept as a
    separate function (not shared) since this report has no min_faltas
    equivalent and the two are likely to diverge further (e.g. a
    minimum-overtime-minutes filter later)."""
    today = datetime.now().date()
    try:
        start_date = datetime.strptime(args.get("start") or today.strftime("%Y-%m-%d"), "%Y-%m-%d").date()
        end_date = datetime.strptime(args.get("end") or today.strftime("%Y-%m-%d"), "%Y-%m-%d").date()
    except ValueError:
        start_date = end_date = today
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    if (end_date - start_date).days > 186:
        end_date = start_date + timedelta(days=186)
    sede = args.get("sede") or None
    return start_date, end_date, sede


def _hours_report_data(db, args):
    """Only days worth a human's attention: a tardiness, actual overtime
    minutes, or a lunch problem. A plain on-time day with no overtime adds
    nothing here, same reasoning as the absences report only listing
    'falta' days instead of every single day in range."""
    start_date, end_date, sede = _hours_report_params(args)
    employees = attendance_rules.employees_in_scope(db, sede=sede)
    attendance = attendance_rules.compute_attendance(db, employees, start_date, end_date)

    rows = []
    for employee, days in attendance.items():
        for d in days:
            notable = (
                d.status == "tardanza"
                or (d.overtime_minutes or 0) > 0
                or d.lunch_status in ("excedido", "no_marcado")
            )
            if not notable:
                continue
            rows.append({
                "employee": employee,
                "date": d.date,
                "schedule_name": d.schedule_name,
                "expected_entry": d.expected_entry,
                "first_punch": d.first_punch,
                "minutes_late": d.minutes_late,
                "expected_exit": d.expected_exit,
                "last_punch": d.last_punch,
                "overtime_minutes": d.overtime_minutes,
                "lunch_status": d.lunch_status,
                "lunch_minutes": d.lunch_minutes,
            })
    rows.sort(key=lambda r: (r["date"], r["employee"].full_name))
    return rows, start_date, end_date, sede


@app.route("/reports/hours")
@module_required("reports_hours")
def report_hours():
    db = SessionLocal()
    try:
        rows, start_date, end_date, sede = _hours_report_data(db, request.args)
        devices = db.query(Device).order_by(Device.name).all()
        sedes = sorted({d.sede for d in devices if d.sede})
        return render_template(
            "report_hours.html",
            rows=rows, sedes=sedes,
            start=start_date.strftime("%Y-%m-%d"), end=end_date.strftime("%Y-%m-%d"),
            sede=sede, lunch_labels=attendance_rules.LUNCH_STATUS_LABELS,
        )
    finally:
        db.close()


@app.route("/reports/hours/export.xlsx")
@module_required("reports_hours")
def report_hours_export():
    db = SessionLocal()
    try:
        rows, start_date, end_date, sede = _hours_report_data(db, request.args)
        headers = [
            "Código", "Empleado", "Fecha", "Horario",
            "Ingreso esperado", "Marcó (ingreso)", "Min. tarde",
            "Salida esperada", "Marcó (salida)", "Min. extra",
            "Refrigerio",
        ]
        table = [[
            r["employee"].code or "", r["employee"].full_name, r["date"].strftime("%Y-%m-%d"), r["schedule_name"] or "",
            r["expected_entry"].strftime("%H:%M") if r["expected_entry"] else "",
            r["first_punch"].strftime("%H:%M:%S") if r["first_punch"] else "",
            r["minutes_late"] or "",
            r["expected_exit"].strftime("%H:%M") if r["expected_exit"] else "",
            r["last_punch"].strftime("%H:%M:%S") if r["last_punch"] else "",
            r["overtime_minutes"] if r["overtime_minutes"] else "",
            attendance_rules.LUNCH_STATUS_LABELS.get(r["lunch_status"], ""),
        ] for r in rows]
        subtitle = f"{start_date.strftime('%Y-%m-%d')} a {end_date.strftime('%Y-%m-%d')}"
        if sede:
            subtitle += f" · Sede: {sede}"
        wb = _styled_workbook("Tardanzas y horas extra", "Reporte de Tardanzas y Horas Extra", subtitle, headers, table)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"reporte_horas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return Response(
            buf.read(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


# ----------------------------------------------------------- permisos ----

@app.route("/permissions", methods=["GET", "POST"])
@super_admin_required
def permissions_page():
    """Lets a super_admin turn a whole module's view routes on/off for
    'viewer' and 'admin' -- see central/auth.py (module_required) for how
    this is enforced and MODULES for the full list. super_admin itself is
    never restricted here, by design."""
    db = SessionLocal()
    try:
        if request.method == "POST":
            changes = []
            existing_rows = {(r.role, r.module): r for r in db.query(RoleModule).all()}
            for role in CONFIGURABLE_ROLES:
                for module in MODULES:
                    enabled = bool(request.form.get(f"{role}__{module}"))
                    default = DEFAULT_MODULE_ACCESS.get(module, {}).get(role, False)
                    existing = existing_rows.get((role, module))
                    if existing:
                        if existing.enabled != enabled:
                            existing.enabled = enabled
                            changes.append(f"{role}/{module}: {'habilitado' if enabled else 'deshabilitado'}")
                    elif enabled != default:
                        db.add(RoleModule(role=role, module=module, enabled=enabled))
                        changes.append(f"{role}/{module}: {'habilitado' if enabled else 'deshabilitado'}")
            db.commit()
            if changes:
                _audit(db, "permissions_edit", detail="; ".join(changes))
                flash("Permisos actualizados.")
            else:
                flash("Sin cambios.")
            return redirect(url_for("permissions_page"))

        current = {
            (role, module): module_enabled(role, module)
            for role in CONFIGURABLE_ROLES for module in MODULES
        }
        return render_template(
            "permissions.html", modules=MODULES, roles=CONFIGURABLE_ROLES,
            role_labels=ROLE_LABELS, current=current,
        )
    finally:
        db.close()


# ------------------------------------------------------- dashboard users ---
# Accounts that can log into this web app (distinct from Employee/device users)

@app.route("/users")
@super_admin_required
def users_list():
    db = SessionLocal()
    try:
        users = db.query(User).order_by(User.username).all()
        return render_template("users.html", users=users)
    finally:
        db.close()


@app.route("/users/new", methods=["GET", "POST"])
@super_admin_required
def user_new():
    if request.method == "POST":
        db = SessionLocal()
        try:
            username = request.form["username"].strip()
            if db.query(User).filter(User.username == username).first():
                flash("Ese usuario ya existe.")
                return redirect(url_for("user_new"))
            role = request.form.get("role", "viewer")
            db.add(User(
                username=username,
                password_hash=hash_password(request.form["password"]),
                role=role,
                active=True,
            ))
            db.commit()
            _audit(db, "account_add", target=username, detail=f"rol: {role}")
            flash(f"Cuenta '{username}' creada.")
            return redirect(url_for("users_list"))
        finally:
            db.close()
    return render_template("user_form.html")


@app.route("/users/<int:user_pk>/toggle", methods=["POST"])
@super_admin_required
def user_toggle(user_pk):
    db = SessionLocal()
    try:
        user = db.get(User, user_pk) or abort_404()
        if user.id == flask_session.get("user_id"):
            flash("No puedes desactivar tu propia cuenta.")
        else:
            user.active = not user.active
            db.commit()
            _audit(db, "account_toggle", target=user.username, detail="activo" if user.active else "desactivado")
        return redirect(url_for("users_list"))
    finally:
        db.close()


@app.route("/users/<int:user_pk>/delete", methods=["POST"])
@super_admin_required
def user_delete(user_pk):
    """Hard delete -- unlike a device (whose history lives in
    attendance_log via a foreign key), there's nothing else in the schema
    that references a User by id except TrustedDevice (cleaned up here) and
    audit_log, which stores the username as a plain string and is
    unaffected by the account going away."""
    db = SessionLocal()
    try:
        user = db.get(User, user_pk) or abort_404()
        if user.id == flask_session.get("user_id"):
            flash("No puedes eliminar tu propia cuenta.")
            return redirect(url_for("users_list"))
        username = user.username
        db.query(TrustedDevice).filter(TrustedDevice.user_id == user.id).delete()
        db.delete(user)
        db.commit()
        _audit(db, "account_delete", target=username)
        flash(f"Cuenta '{username}' eliminada.")
        return redirect(url_for("users_list"))
    finally:
        db.close()


@app.route("/users/<int:user_pk>/reset_2fa", methods=["POST"])
@super_admin_required
def user_reset_2fa(user_pk):
    """Safety valve for a lost-phone lockout: turns 2FA off for that
    account and forgets every trusted device it had, so they can log in
    with just their password and set it up again if they want."""
    db = SessionLocal()
    try:
        user = db.get(User, user_pk) or abort_404()
        user.totp_enabled = False
        user.totp_secret = None
        db.query(TrustedDevice).filter(TrustedDevice.user_id == user.id).delete()
        db.commit()
        _audit(db, "account_reset_2fa", target=user.username)
        flash(f"2FA desactivado para '{user.username}'. Puede iniciar sesión solo con su contraseña.")
        return redirect(url_for("users_list"))
    finally:
        db.close()


@app.route("/users/<int:user_pk>/reset_password", methods=["GET", "POST"])
@super_admin_required
def user_reset_password(user_pk):
    db = SessionLocal()
    try:
        user = db.get(User, user_pk) or abort_404()
        if request.method == "POST":
            new_password = request.form.get("password", "")
            if len(new_password) < 6:
                flash("La contraseña debe tener al menos 6 caracteres.")
                return redirect(url_for("user_reset_password", user_pk=user.id))
            user.password_hash = hash_password(new_password)
            db.commit()
            _audit(db, "account_reset_password", target=user.username)
            flash(f"Contraseña de '{user.username}' actualizada. Avísale la nueva contraseña por un medio seguro.")
            return redirect(url_for("users_list"))
        return render_template("user_reset_password.html", user=user)
    finally:
        db.close()


@app.route("/users/<int:user_pk>/unlock", methods=["POST"])
@super_admin_required
def user_unlock(user_pk):
    """Clears a lockout from MAX_LOGIN_ATTEMPTS failed passwords in a row
    -- the only way back in for that account short of a super_admin doing
    this, by design (see login())."""
    db = SessionLocal()
    try:
        user = db.get(User, user_pk) or abort_404()
        user.locked = False
        user.failed_login_attempts = 0
        db.commit()
        _audit(db, "account_unlock", target=user.username)
        flash(f"Cuenta '{user.username}' desbloqueada.")
        return redirect(url_for("users_list"))
    finally:
        db.close()


# --------------------------------------------------------------- 2FA self --

@app.route("/account/security", methods=["GET"])
@login_required
def account_security():
    db = SessionLocal()
    try:
        user = db.get(User, flask_session["user_id"])
        trusted = db.query(TrustedDevice).filter(TrustedDevice.user_id == user.id).order_by(TrustedDevice.created_at.desc()).all()
        recent_logins = (
            db.query(AuditLog)
            .filter(AuditLog.username == user.username, AuditLog.action.in_(["login", "login_failed", "login_2fa_failed"]))
            .order_by(AuditLog.timestamp.desc()).limit(10).all()
        )
        return render_template(
            "account_security.html", user=user, trusted=trusted, recent_logins=recent_logins,
            session_ip=flask_session.get("client_ip"), session_hostname=flask_session.get("client_hostname"),
        )
    finally:
        db.close()


@app.route("/account/security/change_password", methods=["POST"])
@login_required
def change_password():
    """Self-service, distinct from the super_admin-only /users/<id>/reset_password:
    requires the CURRENT password, so a session left open on someone
    else's machine can't be used to lock the real owner out."""
    db = SessionLocal()
    try:
        user = db.get(User, flask_session["user_id"])
        if not verify_password(request.form.get("current_password", ""), user.password_hash):
            flash("Contraseña actual incorrecta -- no se cambió nada.")
            return redirect(url_for("account_security"))
        new_password = request.form.get("new_password", "")
        if len(new_password) < 6:
            flash("La contraseña nueva debe tener al menos 6 caracteres.")
            return redirect(url_for("account_security"))
        user.password_hash = hash_password(new_password)
        db.commit()
        _audit(db, "account_change_password", target=user.username)
        flash("Contraseña actualizada.")
        return redirect(url_for("account_security"))
    finally:
        db.close()


@app.route("/account/security/2fa/setup", methods=["GET", "POST"])
@login_required
def setup_2fa():
    db = SessionLocal()
    try:
        user = db.get(User, flask_session["user_id"])
        if user.totp_enabled:
            flash("El 2FA ya está activado en tu cuenta.")
            return redirect(url_for("account_security"))

        if request.method == "POST":
            pending_secret = flask_session.get("pending_totp_secret")
            if not pending_secret or not verify_code(pending_secret, request.form.get("code", "")):
                flash("Código incorrecto. Escanea el QR de nuevo si tu app no coincide.")
                return redirect(url_for("setup_2fa"))
            user.totp_secret = pending_secret
            user.totp_enabled = True
            db.commit()
            flask_session.pop("pending_totp_secret", None)
            _audit(db, "2fa_setup", target=user.username)
            flash("2FA activado. La próxima vez que inicies sesión (desde un equipo nuevo) te pedirá el código.")
            return redirect(url_for("account_security"))

        secret = flask_session.get("pending_totp_secret") or generate_secret()
        flask_session["pending_totp_secret"] = secret
        qr = qr_data_uri(secret, user.username)
        return render_template("setup_2fa.html", qr=qr, secret=secret)
    finally:
        db.close()


@app.route("/account/security/2fa/disable", methods=["POST"])
@login_required
def disable_2fa():
    db = SessionLocal()
    try:
        user = db.get(User, flask_session["user_id"])
        if not verify_password(request.form.get("password", ""), user.password_hash):
            flash("Contraseña incorrecta -- no se desactivó el 2FA.")
            return redirect(url_for("account_security"))
        user.totp_enabled = False
        user.totp_secret = None
        db.query(TrustedDevice).filter(TrustedDevice.user_id == user.id).delete()
        db.commit()
        _audit(db, "2fa_disable", target=user.username)
        flash("2FA desactivado en tu cuenta.")
        return redirect(url_for("account_security"))
    finally:
        db.close()


@app.route("/account/security/trusted/<int:trusted_id>/revoke", methods=["POST"])
@login_required
def revoke_trusted_device(trusted_id):
    db = SessionLocal()
    try:
        row = db.get(TrustedDevice, trusted_id)
        if not row or row.user_id != flask_session["user_id"]:
            abort_404()
        db.delete(row)
        db.commit()
        flash("Dispositivo olvidado -- la próxima vez pedirá el código de nuevo ahí.")
        return redirect(url_for("account_security"))
    finally:
        db.close()


# ------------------------------------------------------- backup/restore ---

@app.route("/backup")
@super_admin_required
def backup_page():
    db = SessionLocal()
    try:
        counts = {
            "devices": db.query(Device).count(),
            "employees": db.query(Employee).count(),
            "employee_map": db.query(EmployeeMap).count(),
            "attendance_log": db.query(AttendanceLog).count(),
            "users": db.query(User).count(),
        }
        return render_template("backup.html", counts=counts)
    finally:
        db.close()


@app.route("/backup/export")
@super_admin_required
def backup_export():
    db = SessionLocal()
    try:
        data = export_data(db)
        filename = f"asistencia_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        _audit(db, "backup_export", target=filename)
        return Response(
            json.dumps(data, ensure_ascii=False, indent=2),
            mimetype="application/json",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


@app.route("/backup/import", methods=["POST"])
@super_admin_required
def backup_import():
    if not request.form.get("confirm"):
        flash("Debes marcar la casilla de confirmación para restaurar un backup.")
        return redirect(url_for("backup_page"))
    file = request.files.get("backup_file")
    if not file or not file.filename:
        flash("Selecciona un archivo de backup (.json).")
        return redirect(url_for("backup_page"))
    try:
        data = json.load(file.stream)
    except Exception as e:
        flash(f"El archivo no es un JSON válido: {e}")
        return redirect(url_for("backup_page"))

    db = SessionLocal()
    try:
        counts = import_data(db, data)
        db.commit()
        _audit(db, "backup_import", target=file.filename, detail=", ".join(f"{k}: {v}" for k, v in counts.items()))
        flash(
            "Backup restaurado. "
            + ", ".join(f"{k}: {v}" for k, v in counts.items())
            + ". Tu sesión puede haber quedado inválida si tu cuenta no estaba en el backup — vuelve a iniciar sesión si es necesario."
        )
        return redirect(url_for("backup_page"))
    except Exception as e:
        db.rollback()
        flash(f"Error restaurando el backup (no se aplicó ningún cambio): {e}")
        return redirect(url_for("backup_page"))
    finally:
        db.close()


# ------------------------------------------------------------ audit log ---

@app.route("/audit")
@super_admin_required
def audit_log():
    db = SessionLocal()
    try:
        query = db.query(AuditLog)
        action_filter = request.args.get("action", "").strip()
        username_filter = request.args.get("username", "").strip()
        date_from = request.args.get("from", "").strip()
        date_to = request.args.get("to", "").strip()

        if action_filter:
            query = query.filter(AuditLog.action == action_filter)
        if username_filter:
            query = query.filter(AuditLog.username == username_filter)
        if date_from:
            query = query.filter(AuditLog.timestamp >= datetime.strptime(date_from, "%Y-%m-%d"))
        if date_to:
            query = query.filter(AuditLog.timestamp < datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1))

        page = max(1, request.args.get("page", 1, type=int))
        total = query.count()
        rows = (
            query.order_by(AuditLog.timestamp.desc())
            .offset((page - 1) * PAGE_SIZE).limit(PAGE_SIZE).all()
        )
        actions = [a for (a,) in db.query(AuditLog.action).distinct().order_by(AuditLog.action).all()]
        usernames = [u for (u,) in db.query(AuditLog.username).distinct().filter(AuditLog.username.isnot(None)).order_by(AuditLog.username).all()]
        return render_template(
            "audit.html", rows=rows, total=total, page=page, page_size=PAGE_SIZE,
            actions=actions, usernames=usernames,
            filters={"action": action_filter, "username": username_filter, "from": date_from, "to": date_to},
        )
    finally:
        db.close()


@app.route("/audit/export.csv")
@super_admin_required
def audit_export_csv():
    db = SessionLocal()
    try:
        rows = db.query(AuditLog).order_by(AuditLog.timestamp.desc()).all()
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["Fecha/Hora (UTC)", "Usuario", "Acción", "Objetivo", "Detalle", "IP", "Equipo (PC)", "Navegador"])
        for r in rows:
            writer.writerow([
                r.timestamp.strftime("%Y-%m-%d %H:%M:%S"), r.username or "", r.action,
                r.target or "", r.detail or "", r.ip_address or "", r.hostname or "", r.user_agent or "",
            ])
        filename = f"auditoria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        return Response(
            buf.getvalue(), mimetype="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    finally:
        db.close()


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)
